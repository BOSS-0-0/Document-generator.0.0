import os
import uuid
import datetime
import calendar
import io
import re
import sys
import subprocess
import json
import tempfile
import shutil

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

COMMON_KEY = "_common_"

FIELD_TYPES = ["Текст", "Число", "Дата", "Дата со сдвигом", "Нумерация"]

MONTHS_GENITIVE_RU = [
    "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
]

STANDARD_FIELDS_0 = [
    ("ФИО пациента", "Текст", "{prim0}"),
    ("Дата рождения", "Дата", "{prim1}"),
    ("Пол (муж - 1, жен - 2)", "Текст", "{prim2}"),
]


def new_id():
    """Генерирует и возвращает уникальный идентификатор в формате UUID."""
    return str(uuid.uuid4())


def safe_filename(name):
    """Очищает имя файла от недопустимых в файловой системе символов."""
    name = str(name)
    name = re.sub(r'[\\/*?:"<>|]', "", name)
    name = name.strip()
    return name or "document"


def normalize_key_text(key):
    """Приводит строку ключа к стандартному виду, заключая его в фигурные скобки."""
    key = str(key).strip()
    if not key:
        return key
    if not key.startswith("{"):
        key = "{" + key
    if not key.endswith("}"):
        key += "}"
    return key


def clean_text_piece(text):
    """Удаляет лишние пробелы и символы переноса строки, сжимая текст в одну строку."""
    return " ".join(str(text).split())


def parse_numeration_value(value):
    """Преобразует строковое значение в целое число для полей типа Нумерация."""
    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        raise ValueError("Некорректное число")
    if not number.is_integer():
        raise ValueError("Нужно целое число")
    return int(number)


def parse_date_to_parts(value):
    """Разбивает строку даты в формате ДД.ММ.ГГГГ на отдельные компоненты."""
    value = str(value).strip()
    if not value:
        return "", "", ""
    try:
        dt = datetime.datetime.strptime(value, "%d.%m.%Y")
        return str(dt.day), str(dt.month), str(dt.year)
    except ValueError:
        parts = value.split(".")
        if len(parts) == 3:
            return parts[0].strip(), parts[1].strip(), parts[2].strip()
        return "", "", ""


def russian_month_name(month_number):
    """Возвращает название месяца в родительном падеже по его номеру."""
    if 1 <= month_number <= 12:
        return MONTHS_GENITIVE_RU[month_number - 1]
    return ""


def add_months_to_date(dt, months):
    """Прибавляет указанное количество месяцев к дате, корректно обрабатывая концы месяцев."""
    month_index = dt.month - 1 + months
    year = dt.year + month_index // 12
    month = month_index % 12 + 1
    day = min(dt.day, calendar.monthrange(year, month)[1])
    return dt.replace(year=year, month=month, day=day)


def compute_field_value(field, all_fields=None):
    """Вычисляет итоговое значение поля с учетом его типа и правил сдвига."""
    if field.get("type") == "Дата со сдвигом":
        base = None
        src = field.get("shift_base", "today")
        if src == "today":
            base = datetime.datetime.now()
        elif all_fields:
            for f in all_fields:
                if f.get("id") == src:
                    try:
                        base = datetime.datetime.strptime(
                            str(f.get("value", "")).strip(), "%d.%m.%Y")
                    except ValueError:
                        base = None
                    break
        if base is None:
            return ""
        try:
            months = int(field.get("shift_months", 0))
        except Exception:
            months = 0
        try:
            days = int(field.get("shift_days", 0))
        except Exception:
            days = 0
        dt = add_months_to_date(base, months) + datetime.timedelta(days=days)
        return dt.strftime("%d.%m.%Y")

    value = str(field.get("value", "")).strip()
    if field.get("type") == "Дата":
        try:
            return datetime.datetime.strptime(value, "%d.%m.%Y").strftime("%d.%m.%Y")
        except ValueError:
            return value
    if field.get("type") == "Нумерация":
        try:
            parsed = parse_numeration_value(value)
            return "" if parsed is None else str(parsed)
        except ValueError:
            return value
    return value


def validate_field_value(field):
    """Проверяет корректность значения поля. Возвращает кортеж (успех, сообщение об ошибке)."""
    if field.get("type") == "Дата со сдвигом":
        return True, ""
    value = str(field.get("value", "")).strip()
    if not value:
        return True, ""
    if field.get("type") == "Дата":
        try:
            datetime.datetime.strptime(value, "%d.%m.%Y")
            return True, ""
        except ValueError:
            return False, "Дата должна быть в формате ДД.ММ.ГГГГ, например 01.01.2000"
    if field.get("type") == "Число":
        try:
            float(value.replace(",", ".").replace(" ", ""))
            return True, ""
        except ValueError:
            return False, "Нужно ввести число"
    if field.get("type") == "Нумерация":
        try:
            parse_numeration_value(value)
            return True, ""
        except ValueError:
            return False, "Для типа 'Нумерация' нужно ввести целое число"
    return True, ""


def build_mapping(fields, template=None):
    """Создает словарь замен ключей на значения, включая производные форматы дат."""
    mapping = {}
    existing_keys = {f.get("key") for f in fields if f.get("key")}

    for field in fields:
        key = field.get("key")
        if not key:
            continue
        if field.get("visible", True):
            mapping[key] = compute_field_value(field, fields)
        else:
            mapping[key] = ""

    for field in fields:
        if field.get("type") not in ("Дата", "Дата со сдвигом"):
            continue
        key = field.get("key")
        if not key:
            continue
        value = compute_field_value(field, fields)
        s_day, s_month, s_mnum, s_year = key + "~j", key + "~F", key + "~m", key + "~Y"

        ok_date = None
        if value:
            try:
                ok_date = datetime.datetime.strptime(value, "%d.%m.%Y")
            except ValueError:
                ok_date = None

        if ok_date is None:
            for s in (s_day, s_month, s_mnum, s_year):
                if s not in existing_keys:
                    mapping[s] = ""
            continue

        if s_day not in existing_keys:
            mapping[s_day] = ok_date.strftime("%d")
        if s_month not in existing_keys:
            mapping[s_month] = russian_month_name(ok_date.month)
        if s_mnum not in existing_keys:
            mapping[s_mnum] = ok_date.strftime("%m")
        if s_year not in existing_keys:
            mapping[s_year] = ok_date.strftime("%Y")

    if template is not None:
        mapping.setdefault("{DOC_NAME}", str(template.get("name", "")))
    return mapping


def render_text(text, fields, template=None):
    """Заменяет все ключи в обычном тексте на соответствующие значения из полей."""
    if not text:
        return ""
    mapping = build_mapping(fields, template)
    for key, value in mapping.items():
        text = text.replace(key, str(value))
    return text


def extract_text_from_docx(docx_bytes):
    """Извлекает текстовое содержимое из байтов Word-документа для предпросмотра."""
    doc = Document(io.BytesIO(docx_bytes))
    lines = []
    for paragraph in doc.paragraphs:
        t = clean_text_piece(paragraph.text)
        if t:
            lines.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells, prev = [], None
            for cell in row.cells:
                t = clean_text_piece(cell.text)
                if t and t != prev:
                    cells.append(t)
                prev = t
            line = " | ".join(cells)
            if line:
                lines.append(line)
        lines.append("")
    return "\n".join(lines)


def replace_paragraph(paragraph, mapping):
    """Заменяет ключи в абзаце Word, сохраняя форматирование первого фрагмента (run)."""
    if paragraph.runs:
        full_text = "".join(run.text for run in paragraph.runs)
    else:
        full_text = paragraph.text
    if not full_text:
        return
    new_text = full_text
    for key, value in mapping.items():
        new_text = new_text.replace(key, str(value))
    if new_text == full_text:
        return
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


def replace_in_table(table, mapping):
    """Рекурсивно заменяет ключи во всех ячейках таблицы и вложенных таблицах Word."""
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_paragraph(paragraph, mapping)
            for nested in cell.tables:
                replace_in_table(nested, mapping)


def replace_docx(doc, mapping):
    """Применяет замены ключей ко всему документу Word, включая колонтитулы."""
    for paragraph in doc.paragraphs:
        replace_paragraph(paragraph, mapping)
    for table in doc.tables:
        replace_in_table(table, mapping)
    for section in doc.sections:
        try:
            parts = [section.header, section.footer,
                     section.first_page_header, section.first_page_footer,
                     section.even_page_header, section.even_page_footer]
        except Exception:
            parts = []
        for part in parts:
            if part is None:
                continue
            try:
                for paragraph in part.paragraphs:
                    replace_paragraph(paragraph, mapping)
                for table in part.tables:
                    replace_in_table(table, mapping)
            except Exception:
                pass


def make_docx(template, fields, use_original=False):
    """Генерирует байты Word-документа на основе шаблона и данных полей."""
    mapping = build_mapping(fields, template)
    if use_original and template.get("docx_bytes"):
        doc = Document(io.BytesIO(template["docx_bytes"]))
        replace_docx(doc, mapping)
    else:
        doc = Document()
        rendered = render_text(template.get("text", ""), fields, template)
        for line in rendered.splitlines():
            doc.add_paragraph(line)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def make_xlsx(template, fields):
    """Генерирует байты Excel-файла с таблицей видимых полей и их значений."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    visible = [f for f in fields if f.get("visible", True)]
    if visible:
        ws.append([f["name"] for f in visible])
        ws.append([compute_field_value(f, fields) for f in visible])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        for column in ws.columns:
            max_length = 0
            letter = column[0].column_letter
            for cell in column:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[letter].width = min(max(max_length + 2, 10), 60)
    else:
        ws.append(["Нет видимых полей"])
    info = wb.create_sheet("Инфо")
    info.append(["Параметр", "Значение"])
    info.append(["Название шаблона", template.get("name", "")])
    info.append(["Дата выгрузки", datetime.datetime.now().strftime("%d.%m.%Y %H:%M")])
    for cell in info[1]:
        cell.font = Font(bold=True)
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def open_file(path):
    """Открывает указанный файл программой, ассоциированной с его расширением в ОС."""
    try:
        path = os.path.abspath(path)
        if not os.path.exists(path):
            return False, f"Файл не найден: {path}"
        if sys.platform.startswith("win"):
            try:
                os.startfile(path)
                return True, ""
            except Exception as e:
                try:
                    subprocess.Popen(["cmd", "/c", "start", "", path])
                    return True, ""
                except Exception as e2:
                    return False, f"os.startfile: {e}; cmd start: {e2}"
        if sys.platform == "darwin":
            try:
                subprocess.Popen(["open", path])
                return True, ""
            except Exception as e:
                return False, str(e)
        last_error = "не найдена команда открытия файла"
        for cmd in (["xdg-open", path], ["gio", "open", path],
                    ["gnome-open", path], ["kde-open", path]):
            try:
                proc = subprocess.Popen(cmd, stdout=subprocess.DEVNULL,
                                        stderr=subprocess.PIPE)
                try:
                    _, err = proc.communicate(timeout=2)
                    if proc.returncode == 0:
                        return True, ""
                    last_error = err.decode("utf-8", errors="ignore") or \
                        f"{cmd[0]} вернул код {proc.returncode}"
                except subprocess.TimeoutExpired:
                    return True, ""
            except Exception as e:
                last_error = str(e)
        try:
            import webbrowser
            webbrowser.open(path)
            return True, ""
        except Exception as e:
            return False, f"{last_error}; webbrowser: {e}"
    except Exception as e:
        return False, str(e)


def find_libreoffice():
    """Ищет исполняемый файл LibreOffice в системных путях или стандартных директориях."""
    for name in ("libreoffice", "soffice", "libreoffice.exe", "soffice.exe"):
        path = shutil.which(name)
        if path:
            return path
    if sys.platform.startswith("win"):
        import glob as _glob
        for pattern in (r"C:\Program Files\LibreOffice*\program\soffice.exe",
                        r"C:\Program Files (x86)\LibreOffice*\program\soffice.exe"):
            found = _glob.glob(pattern)
            if found:
                return found[0]
    return None


class DocGeneratorApp:
    """Главный класс приложения: управляет данными, интерфейсом и действиями пользователя."""

    def __init__(self, root):
        """Инициализирует главное окно, переменные состояния и загружает сохраненные данные."""
        self.root = root
        self.root.title("Генератор документации 0-М")
        self.root.geometry("1000x720")

        self.fields = []
        self.field_counter = 0
        self.selected_field_id = None
        self.template_docx_bytes = None

        self.template_text_str = ""
        self.template_source_text = ""

        self.word_app = None
        self.word_doc = None
        self.word_work_path = None

        self.storage_path = os.path.join(
            os.path.expanduser("~"), ".doc_generator_templates.json")
        self.template_states = self._load_storage()
        self.current_template_key = None

        self.field_name_var = tk.StringVar()
        self.field_type_var = tk.StringVar(value=FIELD_TYPES[0])
        self.field_value_var = tk.StringVar()
        self.field_visible_var = tk.BooleanVar(value=True)
        self.field_key_var = tk.StringVar(value="(будет создан автоматически)")

        self.field_date_day_var = tk.StringVar()
        self.field_date_month_var = tk.StringVar()
        self.field_date_year_var = tk.StringVar()

        self.field_shift_base_var = tk.StringVar(value="Сегодня")
        self.field_shift_months_var = tk.StringVar(value="0")
        self.field_shift_days_var = tk.StringVar(value="0")

        self.template_name_var = tk.StringVar(value="Шаблон 0-М")
        self.use_original_var = tk.BooleanVar(value=True)

        self._build_menu()
        self._build_ui()

        default_text = ("Шаблон не открыт. Откройте файл: Файл -> Открыть шаблон.\n"
                        "Кнопка «Поля 0-М» создаст стандартные поля с ключами шаблона.\n")
        self.template_text_str = default_text
        self.template_source_text = default_text

        self._init_fields_from_storage()
        self.current_template_key = self._template_key()
        self._load_template_text(self.current_template_key)

        self.refresh_fields_list()

        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

    def _is_windows(self):
        """Возвращает True, если приложение запущено в операционной системе Windows."""
        return sys.platform.startswith("win")

    def on_close(self):
        """Сохраняет текущее состояние в хранилище и закрывает приложение."""
        self._save_storage()
        self.root.destroy()

    def _load_storage(self):
        """Загружает данные шаблонов и полей из локального JSON-файла."""
        try:
            with open(self.storage_path, "r", encoding="utf-8") as f:
                data = json.load(f)
            if isinstance(data, dict):
                return data
        except Exception:
            pass
        return {}

    def _init_fields_from_storage(self):
        """Инициализирует список общих полей, поддерживая миграцию из старых форматов хранения."""
        if COMMON_KEY not in self.template_states:
            for key, st in self.template_states.items():
                if isinstance(st, dict) and st.get("fields"):
                    self.template_states[COMMON_KEY] = {
                        "fields": st.get("fields", []),
                        "field_counter": st.get("field_counter", 0)}
                    break
        common = self.template_states.setdefault(
            COMMON_KEY, {"fields": [], "field_counter": 0})
        if not isinstance(common.get("fields"), list):
            common["fields"] = []
        self.fields = common["fields"]
        try:
            self.field_counter = int(common.get("field_counter", 0))
        except Exception:
            self.field_counter = 0

    def _save_storage(self):
        """Сохраняет текущие общие поля и тексты шаблонов в локальный JSON-файл."""
        try:
            common = self.template_states.setdefault(
                COMMON_KEY, {"fields": [], "field_counter": 0})
            common["fields"] = self.fields
            common["field_counter"] = self.field_counter

            key = self.current_template_key
            if key:
                tstate = self.template_states.setdefault(key, {})
                tstate["text"] = self.template_text_str
                tstate["source_text"] = self.template_source_text

            with open(self.storage_path, "w", encoding="utf-8") as f:
                json.dump(self.template_states, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    def _template_key(self, name=None):
        """Формирует уникальный строковый ключ для шаблона на основе его названия."""
        if name is None:
            name = self.template_name_var.get()
        name = str(name).strip()
        return name or "Без названия"

    def _load_template_text(self, key):
        """Загружает сохраненный текстовый контент шаблона из хранилища по ключу."""
        state = self.template_states.get(key, {})
        if "text" in state:
            self.template_text_str = state["text"]
            self.template_source_text = state.get("source_text", state["text"])

    def _save_template_text(self, key):
        """Сохраняет текущий текстовый контент шаблона в хранилище под указанным ключом."""
        if not key:
            return
        tstate = self.template_states.setdefault(key, {})
        tstate["text"] = self.template_text_str
        tstate["source_text"] = self.template_source_text

    def sync_template_fields(self, force=False):
        """Синхронизирует данные при смене названия шаблона, сохраняя предыдущее состояние."""
        new_key = self._template_key()
        if self.current_template_key is None:
            self.current_template_key = new_key
            self._load_template_text(new_key)
            return True
        if new_key == self.current_template_key and not force:
            return False
        self._save_template_text(self.current_template_key)
        self.current_template_key = new_key
        self._load_template_text(new_key)
        self._save_storage()
        return True

    def on_template_name_changed(self, event=None):
        """Обработчик события изменения названия шаблона пользователем в интерфейсе."""
        self.sync_template_fields()

    def _build_menu(self):
        """Создает и настраивает верхнее меню приложения с базовыми командами."""
        menubar = tk.Menu(self.root)
        file_menu = tk.Menu(menubar, tearoff=0)
        file_menu.add_command(label="Открыть шаблон .docx/.txt...",
                              command=self.open_template)
        file_menu.add_separator()
        file_menu.add_command(label="Выход", command=self.on_close)
        menubar.add_cascade(label="Файл", menu=file_menu)
        self.root.config(menu=menubar)

    def _build_ui(self):
        """Строит основной графический интерфейс, разделяя его на левую и правую панели."""
        main_pane = ttk.PanedWindow(self.root, orient=tk.HORIZONTAL)
        main_pane.pack(fill=tk.BOTH, expand=True)
        left_frame = ttk.Frame(main_pane, padding=8)
        main_pane.add(left_frame, weight=1)
        right_frame = ttk.Frame(main_pane, padding=8)
        main_pane.add(right_frame, weight=1)

        ttk.Label(left_frame, text="Поля",
                  font=("Segoe UI", 12, "bold")).pack(pady=(0, 6))

        self.field_listbox = tk.Listbox(left_frame, height=10)
        self.field_listbox.pack(fill=tk.BOTH, expand=False)
        self.field_listbox.bind("<<ListboxSelect>>", self.on_field_selected)

        field_buttons = ttk.Frame(left_frame)
        field_buttons.pack(fill=tk.X, pady=5)
        ttk.Button(field_buttons, text="Удалить поле",
                   command=self.delete_field).pack(side=tk.LEFT)
        ttk.Button(field_buttons, text="Вставить данные",
                   command=self.open_paste_dialog).pack(side=tk.LEFT, padx=(5, 0))
        ttk.Button(field_buttons, text="Поля 0-М",
                   command=self.create_standard_fields).pack(side=tk.LEFT, padx=(5, 0))
        ttk.Button(field_buttons, text="Демо",
                   command=self.create_demo).pack(side=tk.RIGHT)

        field_form = ttk.LabelFrame(left_frame, text="Поле", padding=8)
        field_form.pack(fill=tk.BOTH, expand=True, pady=(8, 0))
        field_form.columnconfigure(1, weight=1)

        ttk.Label(field_form, text="Название").grid(row=0, column=0, sticky="w", pady=2)
        ttk.Entry(field_form, textvariable=self.field_name_var).grid(
            row=0, column=1, sticky="ew", padx=(6, 0), pady=2)

        ttk.Label(field_form, text="Тип").grid(row=1, column=0, sticky="w", pady=2)
        self.field_type_combo = ttk.Combobox(
            field_form, textvariable=self.field_type_var,
            values=FIELD_TYPES, state="readonly")
        self.field_type_combo.grid(row=1, column=1, sticky="ew", padx=(6, 0), pady=2)
        self.field_type_combo.bind("<<ComboboxSelected>>", self.on_field_type_changed)

        self.field_value_label = ttk.Label(field_form, text="Значение")
        self.field_value_entry = ttk.Entry(field_form, textvariable=self.field_value_var)
        self.field_value_label.grid(row=2, column=0, sticky="w", pady=2)
        self.field_value_entry.grid(row=2, column=1, sticky="ew", padx=(6, 0), pady=2)

        self.field_date_frame = ttk.Frame(field_form)
        ttk.Label(self.field_date_frame, text="День").grid(row=0, column=0, sticky="w")
        ttk.Entry(self.field_date_frame, textvariable=self.field_date_day_var,
                  width=5).grid(row=0, column=1, padx=(2, 8))
        ttk.Label(self.field_date_frame, text="Месяц").grid(row=0, column=2, sticky="w")
        ttk.Entry(self.field_date_frame, textvariable=self.field_date_month_var,
                  width=5).grid(row=0, column=3, padx=(2, 8))
        ttk.Label(self.field_date_frame, text="Год").grid(row=0, column=4, sticky="w")
        ttk.Entry(self.field_date_frame, textvariable=self.field_date_year_var,
                  width=7).grid(row=0, column=5, padx=(2, 0))
        self.field_date_frame.columnconfigure(6, weight=1)

        self.field_shift_frame = ttk.Frame(field_form)
        ttk.Label(self.field_shift_frame, text="Отсчет от").grid(row=0, column=0, sticky="w")
        self.field_shift_base_combo = ttk.Combobox(
            self.field_shift_frame, textvariable=self.field_shift_base_var,
            state="readonly", width=16)
        self.field_shift_base_combo.grid(row=0, column=1, padx=(2, 8))
        ttk.Label(self.field_shift_frame, text="Месяцев").grid(row=0, column=2, sticky="w")
        ttk.Entry(self.field_shift_frame, textvariable=self.field_shift_months_var,
                  width=5).grid(row=0, column=3, padx=(2, 8))
        ttk.Label(self.field_shift_frame, text="Дней").grid(row=0, column=4, sticky="w")
        ttk.Entry(self.field_shift_frame, textvariable=self.field_shift_days_var,
                  width=5).grid(row=0, column=5)

        ttk.Label(field_form, text="Ключ").grid(row=3, column=0, sticky="w", pady=2)
        ttk.Entry(field_form, textvariable=self.field_key_var).grid(
            row=3, column=1, sticky="ew", padx=(6, 0), pady=2)

        ttk.Checkbutton(field_form, text="Видимое",
                        variable=self.field_visible_var).grid(
            row=4, column=0, columnspan=2, sticky="w", pady=4)

        field_form_buttons = ttk.Frame(field_form)
        field_form_buttons.grid(row=5, column=0, columnspan=2, sticky="ew", pady=(5, 0))
        ttk.Button(field_form_buttons, text="Сохранить поле",
                   command=self.save_field).pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(field_form_buttons, text="Очистить",
                   command=self.clear_field_form).pack(side=tk.LEFT)

        self._update_field_value_visibility()

        ttk.Label(right_frame, text="Шаблон",
                  font=("Segoe UI", 12, "bold")).pack(pady=(0, 6))

        meta_frame = ttk.Frame(right_frame)
        meta_frame.pack(fill=tk.X, pady=3)
        ttk.Label(meta_frame, text="Название").pack(side=tk.LEFT)
        self.template_name_entry = ttk.Entry(
            meta_frame, textvariable=self.template_name_var, width=28)
        self.template_name_entry.pack(side=tk.LEFT, padx=(4, 12))
        self.template_name_entry.bind("<FocusOut>", self.on_template_name_changed)
        self.template_name_entry.bind("<Return>", self.on_template_name_changed)

        hint = ttk.Label(
            right_frame,
            text="Просмотр и правка шаблона — в редакторе:\n"
                 "«Открыть в редакторе» ниже.\n"
                 "Значения подставляются кнопками справа.",
            justify="left")
        hint.pack(anchor="w", pady=(0, 6))

        export_frame = ttk.Frame(right_frame)
        export_frame.pack(fill=tk.X, pady=(8, 0))
        ttk.Checkbutton(export_frame,
                        text="Использовать исходный Word-шаблон, если он загружен",
                        variable=self.use_original_var).pack(side=tk.LEFT)
        ttk.Button(export_frame, text="Скачать Excel",
                   command=self.download_excel).pack(side=tk.RIGHT, padx=(5, 0))
        ttk.Button(export_frame, text="Скачать Word",
                   command=self.download_word).pack(side=tk.RIGHT)

        word_frame = ttk.LabelFrame(
            right_frame, text="Прямая работа с документом (Word / LibreOffice)", padding=6)
        word_frame.pack(fill=tk.X, pady=(8, 0))
        ttk.Button(word_frame, text="Открыть в редакторе",
                   command=self.open_in_editor).pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(word_frame, text="Вставить данные в документ",
                   command=self.insert_data_into_editor).pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(word_frame, text="Прочитать из редактора",
                   command=self.read_from_editor).pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(word_frame, text="Сохранить документ...",
                   command=self.save_editor_doc).pack(side=tk.LEFT)

    def _refresh_shift_base_choices(self):
        """Обновляет список доступных базовых дат для поля 'Дата со сдвигом'."""
        names = ["Сегодня"]
        for f in self.fields:
            if f.get("type") == "Дата":
                names.append(f.get("name", ""))
        cur = self.field_shift_base_var.get()
        if cur and cur not in names:
            names.append(cur)
        self.field_shift_base_combo["values"] = names

    def _update_field_value_visibility(self):
        """Переключает видимость элементов ввода в форме в зависимости от выбранного типа поля."""
        t = self.field_type_var.get()
        if t == "Дата":
            self.field_value_label.grid_remove()
            self.field_value_entry.grid_remove()
            self.field_shift_frame.grid_remove()
            self.field_date_frame.grid(row=2, column=0, columnspan=2, sticky="ew", pady=2)
        elif t == "Дата со сдвигом":
            self.field_value_label.grid_remove()
            self.field_value_entry.grid_remove()
            self.field_date_frame.grid_remove()
            self._refresh_shift_base_choices()
            self.field_shift_frame.grid(row=2, column=0, columnspan=2, sticky="ew", pady=2)
        else:
            self.field_date_frame.grid_remove()
            self.field_shift_frame.grid_remove()
            self.field_value_label.grid(row=2, column=0, sticky="w", pady=2)
            self.field_value_entry.grid(row=2, column=1, sticky="ew", padx=(6, 0), pady=2)

    def on_field_type_changed(self, event=None):
        """Обрабатывает событие изменения типа поля в выпадающем списке интерфейса."""
        self._update_field_value_visibility()

    def compose_date_value(self):
        """Собирает и валидирует дату из трех отдельных строковых компонентов."""
        day = self.field_date_day_var.get().strip()
        month = self.field_date_month_var.get().strip()
        year = self.field_date_year_var.get().strip()
        if not day and not month and not year:
            return "", True, ""
        if not (day and month and year):
            return "", False, "Введите день, месяц и год полностью."
        try:
            d, m, y = int(day), int(month), int(year)
        except ValueError:
            return "", False, "День, месяц и год должны быть числами."
        if not (1 <= d <= 31 and 1 <= m <= 12 and 1 <= y <= 9999):
            return "", False, "Некорректные компоненты даты."
        value = f"{d:02d}.{m:02d}.{y:04d}"
        valid, message = validate_field_value({"type": "Дата", "value": value})
        if not valid:
            return "", False, message
        return value, True, ""

    def _key_exists(self, key, exclude_id=None):
        """Проверяет, не занят ли указанный ключ другим полем в списке."""
        for f in self.fields:
            if f.get("id") == exclude_id:
                continue
            if f.get("key") == key:
                return True
        return False

    def generate_field_key(self):
        """Генерирует новый уникальный ключ для поля, избегая дубликатов."""
        used = {f.get("key") for f in self.fields if f.get("key")}
        while True:
            index = self.field_counter
            key = f"{{Xkey{index}-{index + 1}X}}"
            self.field_counter += 1
            if key not in used:
                return key

    def add_field_object(self, name, field_type, value, visible, key=None):
        """Создает и возвращает словарь с данными нового поля."""
        field = {"id": new_id(), "name": name,
                 "key": key if key else self.generate_field_key(),
                 "type": field_type, "value": value, "visible": visible}
        self.fields.append(field)
        return field

    def create_standard_fields(self):
        """Массово создает стандартные поля формы 0-М, пропуская уже существующие ключи."""
        added = 0
        for name, ftype, key in STANDARD_FIELDS_0:
            if self._key_exists(key):
                continue
            self.add_field_object(name, ftype, "", True, key=key)
            added += 1
        if added:
            self.refresh_fields_list()
            self._save_storage()
            messagebox.showinfo("Поля 0-М", f"Создано стандартных полей: {added}.")
        else:
            messagebox.showinfo("Поля 0-М", "Стандартные поля уже созданы.")

    def refresh_fields_list(self, select_id=None):
        """Перерисовывает список полей в интерфейсе и обновляет выделение."""
        self.field_listbox.delete(0, tk.END)
        select_index = None
        for index, field in enumerate(self.fields):
            vis = "видимое" if field.get("visible", True) else "скрыто"
            extra = ""
            if field.get("type") == "Дата со сдвигом":
                extra = " | = " + compute_field_value(field, self.fields)
            self.field_listbox.insert(
                tk.END, f"{field.get('name', '')} | {field.get('key', '')} | {vis}{extra}")
            if select_id == field.get("id"):
                select_index = index
            elif select_id is None and self.selected_field_id == field.get("id"):
                select_index = index
        if self.fields:
            if select_index is None:
                select_index = 0
            self.field_listbox.selection_clear(0, tk.END)
            self.field_listbox.selection_set(select_index)
            self.load_field_form(select_index)
        else:
            self.clear_field_form()

    def load_field_form(self, index):
        """Загружает данные выбранного поля из списка в форму редактирования."""
        if 0 <= index < len(self.fields):
            field = self.fields[index]
            self.selected_field_id = field.get("id")
            self.field_name_var.set(field.get("name", ""))
            self.field_type_var.set(field.get("type", FIELD_TYPES[0]))
            self.field_visible_var.set(field.get("visible", True))
            self.field_key_var.set(field.get("key", ""))
            value = str(field.get("value", ""))
            self.field_value_var.set(value)

            if field.get("type") == "Дата":
                d, m, y = parse_date_to_parts(value)
                self.field_date_day_var.set(d)
                self.field_date_month_var.set(m)
                self.field_date_year_var.set(y)
            else:
                self.field_date_day_var.set("")
                self.field_date_month_var.set("")
                self.field_date_year_var.set("")

            if field.get("type") == "Дата со сдвигом":
                self.field_shift_months_var.set(str(field.get("shift_months", 0)))
                self.field_shift_days_var.set(str(field.get("shift_days", 0)))
                base_id = field.get("shift_base", "today")
                base_name = "Сегодня"
                if base_id != "today":
                    for f in self.fields:
                        if f.get("id") == base_id:
                            base_name = f.get("name", "Сегодня")
                            break
                self.field_shift_base_var.set(base_name)
            else:
                self.field_shift_base_var.set("Сегодня")
                self.field_shift_months_var.set("0")
                self.field_shift_days_var.set("0")

            self._update_field_value_visibility()
        else:
            self.clear_field_form()

    def on_field_selected(self, event):
        """Обрабатывает событие выбора элемента в списке полей."""
        selection = self.field_listbox.curselection()
        if not selection:
            return
        self.load_field_form(selection[0])

    def get_selected_field(self):
        """Возвращает словарь данных текущего выбранного поля или None."""
        selection = self.field_listbox.curselection()
        if selection:
            return self.fields[selection[0]]
        if self.selected_field_id:
            for f in self.fields:
                if f.get("id") == self.selected_field_id:
                    return f
        return None

    def clear_field_form(self):
        """Очищает все поля формы редактирования и сбрасывает выделение."""
        self.selected_field_id = None
        self.field_name_var.set("")
        self.field_type_var.set(FIELD_TYPES[0])
        self.field_value_var.set("")
        self.field_visible_var.set(True)
        self.field_key_var.set("(будет создан автоматически)")
        self.field_date_day_var.set("")
        self.field_date_month_var.set("")
        self.field_date_year_var.set("")
        self.field_shift_base_var.set("Сегодня")
        self.field_shift_months_var.set("0")
        self.field_shift_days_var.set("0")
        if hasattr(self, "field_listbox"):
            self.field_listbox.selection_clear(0, tk.END)
        self._update_field_value_visibility()

    def save_field(self):
        """Сохраняет новое или обновляет существующее поле с полной валидацией данных."""
        name = self.field_name_var.get().strip()
        if not name:
            messagebox.showerror("Ошибка", "Укажите название поля")
            return

        field_type = self.field_type_var.get()
        visible = self.field_visible_var.get()
        old_field = self.get_selected_field()

        placeholder = "(будет создан автоматически)"
        key_text = self.field_key_var.get().strip()
        if key_text in ("", placeholder):
            key_text = None
        else:
            key_text = normalize_key_text(key_text)
        if key_text is not None:
            exclude = old_field.get("id") if old_field else None
            if self._key_exists(key_text, exclude_id=exclude):
                messagebox.showerror("Ошибка",
                                     f"Ключ {key_text} уже используется другим полем")
                return

        shift_params = None
        if field_type == "Дата со сдвигом":
            try:
                months = int(self.field_shift_months_var.get().strip() or 0)
            except ValueError:
                messagebox.showwarning("Проверка", "Месяцы должны быть целым числом.")
                return
            try:
                days = int(self.field_shift_days_var.get().strip() or 0)
            except ValueError:
                messagebox.showwarning("Проверка", "Дни должны быть целым числом.")
                return
            base_name = self.field_shift_base_var.get().strip() or "Сегодня"
            base_id = "today"
            if base_name != "Сегодня":
                for f in self.fields:
                    if f.get("name") == base_name and f.get("type") == "Дата":
                        base_id = f.get("id")
                        break
            shift_params = {"shift_base": base_id,
                            "shift_months": months, "shift_days": days}

        if field_type == "Дата":
            value, ok, message = self.compose_date_value()
            if not ok:
                messagebox.showwarning("Проверка даты", message)
                return
        elif field_type == "Дата со сдвигом":
            value = ""
        elif field_type == "Нумерация":
            raw = self.field_value_var.get().strip()
            if raw:
                try:
                    parsed = parse_numeration_value(raw)
                except ValueError:
                    messagebox.showwarning("Проверка значения",
                                           "Для типа 'Нумерация' нужно ввести целое число.")
                    return
                should = False
                if old_field is None:
                    should = True
                else:
                    old_v = str(old_field.get("value", "")).strip()
                    if old_v != raw or old_field.get("type") != field_type:
                        should = True
                value = str(parsed + 1) if should else str(parsed)
            else:
                value = ""
        else:
            value = self.field_value_var.get().strip()
            valid, message = validate_field_value({"type": field_type, "value": value})
            if not valid:
                messagebox.showwarning("Проверка значения", message)
                return

        if old_field:
            old_field["name"] = name
            old_field["type"] = field_type
            old_field["value"] = value
            old_field["visible"] = visible
            if key_text:
                old_field["key"] = key_text
            if shift_params:
                old_field.update(shift_params)
            select_id = old_field["id"]
        else:
            if key_text is None:
                key_text = self.generate_field_key()
            field = self.add_field_object(name, field_type, value, visible, key=key_text)
            if shift_params:
                field.update(shift_params)
            select_id = field["id"]

        self.refresh_fields_list(select_id=select_id)
        self._save_storage()

    def delete_field(self):
        """Удаляет выбранное поле из списка после подтверждения пользователем."""
        field = self.get_selected_field()
        if not field:
            messagebox.showinfo("Удаление поля", "Сначала выберите поле")
            return
        if not messagebox.askyesno("Удаление поля",
                                   f"Удалить поле '{field.get('name', '')}'?"):
            return
        self.fields[:] = [f for f in self.fields if f.get("id") != field.get("id")]
        self.selected_field_id = None
        self.refresh_fields_list()
        self._save_storage()

    def _get_clipboard_text(self):
        """Пытается прочитать текстовые данные из системного буфера обмена."""
        for clip_type in (None, "UTF8_STRING", "STRING", "TEXT"):
            try:
                data = self.root.clipboard_get() if clip_type is None \
                    else self.root.clipboard_get(type=clip_type)
                if isinstance(data, bytes):
                    data = data.decode("utf-8", errors="ignore")
                data = str(data)
                if data.strip():
                    return data
            except Exception:
                continue
        return ""

    def _paste_into_text(self, text_widget):
        """Вставляет содержимое буфера обмена в указанный текстовый виджет."""
        clip = self._get_clipboard_text()
        if clip:
            text_widget.insert(tk.INSERT, clip)
        else:
            messagebox.showwarning("Буфер обмена",
                                   "Буфер обмена пуст или содержит не текст.")

    def _select_all_text(self, text_widget):
        """Выделяет весь текст в указанном текстовом виджете."""
        text_widget.tag_add("sel", "1.0", "end")

    def open_paste_dialog(self):
        """Открывает модальное окно для массовой вставки и распознавания данных."""
        win = tk.Toplevel(self.root)
        win.title("Вставка данных")
        win.geometry("760x480")
        win.transient(self.root)

        hint = ("Вставьте данные, каждое значение на отдельной строке.\n"
                "Понимаются строки вида:\n"
                "   ФИО Никита\n   ФИО: Никита\n   ФИО = Никита\n"
                "   {UfCrm1702382207} Никита\n"
                "Кнопка «Вставить в документ» — применяет данные и сразу\n"
                "собирает/открывает готовый документ (минимум действий).")
        ttk.Label(win, text=hint, justify="left").pack(anchor="w", padx=10, pady=(10, 5))

        buttons = ttk.Frame(win)
        buttons.pack(side=tk.BOTTOM, fill=tk.X, padx=10, pady=(0, 10))
        text = tk.Text(win, wrap="word")

        def paste_and_go():
            if self.paste_and_fill_document(text.get("1.0", "end-1c")):
                win.destroy()

        ttk.Button(buttons, text="Вставить в документ",
                   command=paste_and_go).pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(buttons, text="Вставить из буфера обмена",
                   command=lambda: self._paste_into_text(text)).pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(buttons, text="Обновить значения",
                   command=lambda: self._apply_paste(win, text.get("1.0", "end-1c"))).pack(
            side=tk.LEFT, padx=(0, 5))
        ttk.Button(buttons, text="Закрыть", command=win.destroy).pack(side=tk.LEFT)

        text.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)

        popup = tk.Menu(win, tearoff=0)
        popup.add_command(label="Вставить", command=lambda: self._paste_into_text(text))
        popup.add_command(label="Выделить всё", command=lambda: self._select_all_text(text))
        popup.add_command(label="Очистить", command=lambda: text.delete("1.0", tk.END))

        def on_right_click(event):
            try:
                popup.tk_popup(event.x_root, event.y_root)
            finally:
                popup.grab_release()

        text.bind("<Button-3>", on_right_click)

        clip = self._get_clipboard_text()
        if clip:
            text.insert("1.0", clip)
        text.focus_set()

    def _match_line_to_field(self, line):
        """Пытается сопоставить строку текста с именем или ключом существующего поля."""
        line = line.replace("\xa0", " ")
        line_lower = line.lower()

        for field in self.fields:
            if field.get("type") == "Дата со сдвигом":
                continue
            key = str(field.get("key", "")).strip()
            if not key:
                continue
            variants = [key]
            inner = key[1:-1] if key.startswith("{") and key.endswith("}") else key
            if inner and inner != key:
                variants.append(inner)
            for variant in variants:
                if variant and variant in line:
                    value = line.split(variant, 1)[1]
                    value = value.strip().strip(":=-–—| \t").strip()
                    if value:
                        return field, value

        for field in self.fields:
            if field.get("type") == "Дата со сдвигом":
                continue
            name = str(field.get("name", "")).strip()
            if not name:
                continue
            if line_lower.startswith(name.lower()):
                value = line[len(name):]
                value = value.strip().strip(":=-–—| \t").strip()
                if value:
                    return field, value

        for sep in ("\t", ":", "="):
            if sep not in line:
                continue
            label, _, rest = line.partition(sep)
            label = label.strip().lower()
            rest = rest.strip()
            for field in self.fields:
                if field.get("type") == "Дата со сдвигом":
                    continue
                name = str(field.get("name", "")).strip()
                if name and name.lower() == label and rest:
                    return field, rest
        return None, None

    def apply_pasted_text(self, text):
        """Применяет распознанные данные из текста к соответствующим полям шаблона."""
        updated, unchanged, not_matched = [], [], []
        for raw_line in text.splitlines():
            line = raw_line.strip().replace("\xa0", " ")
            if not line:
                continue
            field, value = self._match_line_to_field(line)
            if field is None or not value:
                not_matched.append(line)
                continue
            if field.get("type") == "Дата":
                value = value.replace(" г.", "").replace("г.", "").strip()
            old = str(field.get("value", ""))
            if old != value:
                field["value"] = value
                updated.append((field.get("name", ""), old, value))
            else:
                unchanged.append(field.get("name", ""))
        return updated, unchanged, not_matched

    def _apply_paste(self, win, pasted):
        """Применяет вставленные данные и показывает пользователю отчет о результатах."""
        if not pasted.strip():
            messagebox.showinfo("Вставка данных",
                                "Нечего применять: вставьте текст с данными.", parent=win)
            return
        updated, unchanged, not_matched = self.apply_pasted_text(pasted)
        self.refresh_fields_list()
        self._save_storage()

        report = []
        if updated:
            report.append("Изменено:\n" + "\n".join(
                f"- {n}: было «{o}» → стало «{v}»" for n, o, v in updated))
        if unchanged:
            report.append("Остались прежними:\n" + "\n".join(f"- {n}" for n in unchanged))
        if not_matched:
            report.append("Не распознано:\n" + "\n".join(f"- {s}" for s in not_matched))
        if not updated:
            report.append("НИЧЕГО НЕ ИЗМЕНИЛОСЬ")
        messagebox.showinfo("Вставка данных", "\n\n".join(report), parent=win)

    def _replace_in_open_word(self, mapping):
        """Выполняет поиск и замену ключей в активном документе Word через COM-интерфейс."""
        replaced = 0
        for key, value in mapping.items():
            if not key:
                continue
            try:
                rng = self.word_doc.Content
                fr = rng.Find
                fr.ClearFormatting()
                fr.Replacement.ClearFormatting()
                if fr.Execute(FindText=str(key), ReplaceWith=str(value),
                              Replace=2, Wrap=1, Forward=True):
                    replaced += 1
            except Exception:
                pass
        return replaced

    def paste_and_fill_document(self, pasted):
        """Комплексно применяет вставленные данные и сразу открывает готовый документ."""
        if not pasted.strip():
            messagebox.showinfo("Вставка данных", "Нечего применять: вставьте текст с данными.")
            return False

        updated, unchanged, not_matched = self.apply_pasted_text(pasted)
        self.refresh_fields_list()
        self._save_storage()

        mapping = build_mapping(self.fields, self.template_data())

        if self._is_windows() and self.word_doc is not None:
            replaced = self._replace_in_open_word(mapping)
            note = f"\nНе распознано строк: {len(not_matched)}." if not_matched else ""
            messagebox.showinfo(
                "Готово",
                f"Полей обновлено: {len(updated)}.\nЗамен в Word: {replaced}."
                f"\nДокумент с данными открыт в Word.{note}")
            return True

        if not self.template_docx_bytes:
            messagebox.showwarning("Внимание",
                                   "Сначала откройте шаблон .docx: Файл -> Открыть шаблон.")
            return False

        stamp = datetime.datetime.now().strftime("%H%M%S")
        work_path = os.path.join(
            tempfile.gettempdir(),
            safe_filename(self.template_name_var.get()) + f"_рабочая_{stamp}.docx")

        try:
            doc = Document(io.BytesIO(self.template_docx_bytes))
            replace_docx(doc, mapping)
            doc.save(work_path)
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось подготовить документ:\n{e}")
            return False

        self.word_work_path = work_path

        try:
            opened, err = self._open_working_copy(work_path)
        except Exception as e:
            opened, err = False, str(e)

        note = f"\nНе распознано строк: {len(not_matched)}." if not_matched else ""
        if not opened:
            messagebox.showwarning(
                "Готово",
                f"Документ с данными подготовлен:\n{work_path}\n\n"
                f"Не удалось открыть автоматически:\n{err}{note}")
            return True

        messagebox.showinfo(
            "Готово",
            f"Полей обновлено: {len(updated)}.\nДокумент с данными открыт в редакторе.{note}")
        return True

    def create_demo(self):
        """Создает набор демонстрационных полей для тестирования функционала программы."""
        if self.fields:
            if not messagebox.askyesno("Демо", "Добавить демо-поля?"):
                return
        f1 = self.add_field_object("ФИО", "Текст", "Иванов Иван", True)
        self.add_field_object("Возраст", "Число", "42", True)
        self.add_field_object("Адрес", "Текст", "Городская ул. 4 дом. 5", True)
        f4 = self.add_field_object("Дата рождения", "Дата", "20.12.1987", True)
        f5 = self.add_field_object("Дата через 5 мес 10 дн", "Дата со сдвигом", "", True)
        f5.update({"shift_base": f4["id"], "shift_months": 5, "shift_days": 10})
        self.refresh_fields_list(select_id=f1["id"])
        self._save_storage()

    def template_data(self):
        """Возвращает словарь с текущими метаданными и содержимым шаблона."""
        return {"name": self.template_name_var.get().strip(),
                "text": self.template_text_str,
                "docx_bytes": self.template_docx_bytes}

    def open_template(self):
        """Открывает диалог выбора файла и загружает содержимое шаблона Word или TXT."""
        path = filedialog.askopenfilename(
            title="Открыть шаблон",
            filetypes=[("Шаблоны", "*.docx *.txt"), ("Word", "*.docx"),
                       ("Текст", "*.txt"), ("Все файлы", "*.*")])
        if not path:
            return
        name = os.path.basename(path)
        try:
            if path.lower().endswith(".docx"):
                with open(path, "rb") as f:
                    content = f.read()
                text = extract_text_from_docx(content)
                self.template_docx_bytes = content
            else:
                try:
                    with open(path, "r", encoding="utf-8") as f:
                        text = f.read()
                except UnicodeDecodeError:
                    with open(path, "r", encoding="cp1251", errors="ignore") as f:
                        text = f.read()
                self.template_docx_bytes = None

            self.template_name_var.set(os.path.splitext(name)[0])
            self.sync_template_fields()
            self.template_text_str = text
            self.template_source_text = text
            self._save_storage()
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось открыть файл:\n{e}")

    def confirm_fields_valid(self):
        """Проверяет все поля на корректность перед началом процесса экспорта."""
        errors = []
        for field in self.fields:
            valid, message = validate_field_value(field)
            if not valid:
                errors.append(f"{field.get('name', '')}: {message}")
        if errors:
            return messagebox.askyesno("Найдены ошибки в полях",
                                       "\n".join(errors) + "\n\nПродолжить экспорт?")
        return True

    def _increment_numeration_after_export(self):
        """Автоматически увеличивает значение всех полей типа 'Нумерация' после выгрузки."""
        changed = False
        for field in self.fields:
            if field.get("type") != "Нумерация":
                continue
            value = str(field.get("value", "")).strip()
            if not value:
                continue
            try:
                field["value"] = str(parse_numeration_value(value) + 1)
                changed = True
            except ValueError:
                pass
        if changed:
            self.refresh_fields_list()
            self._save_storage()

    def _get_word_app(self):
        """Подключается к запущенному экземпляру Microsoft Word через COM-интерфейс."""
        if getattr(self, "word_app", None) is not None:
            return self.word_app
        try:
            import win32com.client as wc
            try:
                app = wc.GetActiveObject("Word.Application")
            except Exception:
                app = wc.Dispatch("Word.Application")
            app.Visible = True
            self.word_app = app
            return app
        except Exception:
            pass
        try:
            import comtypes.client as cc
            try:
                app = cc.GetActiveObject("Word.Application")
            except Exception:
                app = cc.CreateObject("Word.Application")
            app.Visible = True
            self.word_app = app
            return app
        except Exception as e:
            raise RuntimeError(
                "Не удалось подключиться к Microsoft Word.\n"
                "Нужен установленный Word и одна из библиотек:\n"
                "   pip install pywin32\nили\n   pip install comtypes\n"
                f"Ошибка: {e}")

    def _open_working_copy(self, path):
        """Открывает временную копию документа в Word или LibreOffice в зависимости от ОС."""
        if self._is_windows():
            try:
                app = self._get_word_app()
                self.word_doc = app.Documents.Open(os.path.abspath(path))
                return True, ""
            except Exception:
                self.word_doc = None
        lo = find_libreoffice()
        if lo:
            try:
                subprocess.Popen([lo, "--writer", os.path.abspath(path)])
                return True, ""
            except Exception as e:
                last = str(e)
        else:
            last = "LibreOffice не найден"
        opened, err = open_file(path)
        if opened:
            return True, ""
        return False, (f"{err or last}\n\nДля работы с редактором установите "
                       "Microsoft Word или LibreOffice (https://www.libreoffice.org).")

    def open_in_editor(self):
        """Открывает текущий шаблон во внешнем редакторе для ручного редактирования."""
        template = self.template_data()
        if not template.get("docx_bytes"):
            messagebox.showwarning("Редактор",
                                   "Сначала откройте шаблон .docx: Файл -> Открыть шаблон.")
            return
        if self._is_windows() and self.word_doc is not None:
            try:
                self.word_doc.Activate()
                return
            except Exception:
                self.word_doc = None

        work_path = os.path.join(tempfile.gettempdir(),
                                 safe_filename(template["name"]) + "_working.docx")
        with open(work_path, "wb") as f:
            f.write(template["docx_bytes"])
        self.word_work_path = work_path

        try:
            opened, err = self._open_working_copy(work_path)
        except Exception as e:
            opened, err = False, str(e)
        if not opened:
            messagebox.showerror("Редактор", f"Не удалось открыть документ в редакторе:\n{err}")
            return
        if self.word_doc is not None:
            messagebox.showinfo("Редактор",
                                "Документ открыт в Microsoft Word.\n"
                                "Правьте его прямо там, программа вставит данные на живую.")
        else:
            messagebox.showinfo("Редактор",
                                "Документ открыт в редакторе (LibreOffice).\n"
                                "Правьте и сохраняйте его там.\n"
                                "Программа вставит данные в файл и прочитает его обратно.")

    def insert_data_into_editor(self):
        """Вставляет актуальные значения полей в файл, открытый во внешнем редакторе."""
        self.sync_template_fields()
        mapping = build_mapping(self.fields, self.template_data())

        if self._is_windows() and self.word_doc is not None:
            replaced = self._replace_in_open_word(mapping)
            messagebox.showinfo("Редактор",
                                f"Данные вставлены в открытый Word.\nЗамен сделано: {replaced}")
            return

        path = getattr(self, "word_work_path", None)
        if not path or not os.path.exists(path):
            messagebox.showinfo("Редактор", "Сначала нажмите «Открыть в редакторе».")
            return
        if not messagebox.askyesno("Внимание",
                                   "Чтобы вставить данные в файл, ЗАКРОЙТЕ документ "
                                   "в редакторе.\n\nПродолжить?"):
            return
        try:
            with open(path, "rb") as f:
                doc = Document(io.BytesIO(f.read()))
            parts = [p.text for p in doc.paragraphs]
            for t in doc.tables:
                for row in t.rows:
                    for c in row.cells:
                        parts.append(c.text)
            before = "\n".join(parts)
            found = sum(before.count(k) for k in mapping if k)
            replace_docx(doc, mapping)
            doc.save(path)
        except Exception as e:
            messagebox.showerror("Редактор",
                                 f"Не удалось вставить данные:\n{e}\n\n"
                                 "Возможно, документ все еще открыт в редакторе.")
            return
        if messagebox.askyesno("Готово",
                               f"Данные вставлены в файл ({found} замен).\n\n"
                               "Открыть документ в редакторе снова?"):
            self._open_working_copy(path)

    def read_from_editor(self):
        """Читает изменения, сделанные во внешнем редакторе, обратно в память программы."""
        path = getattr(self, "word_work_path", None)

        if self._is_windows() and self.word_doc is not None:
            try:
                self.word_doc.Save()
            except Exception:
                pass

        if not path or not os.path.exists(path):
            messagebox.showinfo("Редактор", "Сначала нажмите «Открыть в редакторе».")
            return

        try:
            with open(path, "rb") as f:
                content = f.read()
            self.template_docx_bytes = content
            self.template_text_str = extract_text_from_docx(content)
            self.template_source_text = self.template_text_str
        except Exception as e:
            messagebox.showerror("Редактор", f"Не удалось прочитать файл:\n{e}")
            return

        self._save_storage()
        messagebox.showinfo("Редактор", "Шаблон обновлен из редактора.")

    def save_editor_doc(self):
        """Сохраняет итоговую версию документа из редактора на диск пользователя."""
        path = filedialog.asksaveasfilename(
            title="Сохранить документ", defaultextension=".docx",
            initialfile=f"{safe_filename(self.template_name_var.get())}.docx",
            filetypes=[("Word документ", "*.docx")])
        if not path:
            return
        try:
            if self._is_windows() and self.word_doc is not None:
                self.word_doc.SaveAs(os.path.abspath(path))
            else:
                work = getattr(self, "word_work_path", None)
                if not work or not os.path.exists(work):
                    messagebox.showinfo("Редактор", "Сначала нажмите «Открыть в редакторе».")
                    return
                shutil.copyfile(work, path)
        except Exception as e:
            messagebox.showerror("Редактор", f"Не удалось сохранить файл:\n{e}")
            return
        opened, open_error = open_file(path)
        self._increment_numeration_after_export()
        if opened:
            messagebox.showinfo("Экспорт", f"Документ сохранен и открыт:\n{path}")
        else:
            messagebox.showwarning("Экспорт",
                                   f"Документ сохранен:\n{path}\n\n"
                                   f"Не удалось открыть автоматически.\nОшибка: {open_error}")

    def download_word(self):
        """Генерирует, сохраняет на диск и открывает Word-файл с подставленными данными."""
        self.sync_template_fields()
        if not self.confirm_fields_valid():
            return
        template = self.template_data()
        use_original = self.use_original_var.get() and bool(template.get("docx_bytes"))
        try:
            data = make_docx(template, self.fields, use_original=use_original)
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось создать Word:\n{e}")
            return
        path = filedialog.asksaveasfilename(
            title="Сохранить Word", defaultextension=".docx",
            initialfile=f"{safe_filename(template['name'])}.docx",
            filetypes=[("Word документ", "*.docx")])
        if not path:
            return
        if not path.lower().endswith(".docx"):
            path += ".docx"
        try:
            with open(path, "wb") as f:
                f.write(data)
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить файл:\n{e}")
            return
        opened, open_error = open_file(path)
        self._increment_numeration_after_export()
        if opened:
            messagebox.showinfo("Экспорт", f"Word файл сохранен и открыт:\n{path}")
        else:
            messagebox.showwarning("Экспорт",
                                   f"Word файл сохранен:\n{path}\n\n"
                                   f"Не удалось открыть автоматически.\nОшибка: {open_error}")

    def download_excel(self):
        """Генерирует, сохраняет на диск и открывает Excel-файл с таблицей данных полей."""
        self.sync_template_fields()
        if not self.confirm_fields_valid():
            return
        template = self.template_data()
        try:
            data = make_xlsx(template, self.fields)
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось создать Excel:\n{e}")
            return
        path = filedialog.asksaveasfilename(
            title="Сохранить Excel", defaultextension=".xlsx",
            initialfile=f"{safe_filename(template['name'])}.xlsx",
            filetypes=[("Excel файл", "*.xlsx")])
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        try:
            with open(path, "wb") as f:
                f.write(data)
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить файл:\n{e}")
            return
        opened, open_error = open_file(path)
        self._increment_numeration_after_export()
        if opened:
            messagebox.showinfo("Экспорт", f"Excel файл сохранен и открыт:\n{path}")
        else:
            messagebox.showwarning("Экспорт",
                                   f"Excel файл сохранен:\n{path}\n\n"
                                   f"Не удалось открыть автоматически.\nОшибка: {open_error}")


def main():
    """Точка входа в приложение: создает главное окно и запускает цикл событий."""
    root = tk.Tk()
    DocGeneratorApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()